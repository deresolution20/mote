from __future__ import annotations

import argparse
import json
import platform
import re
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


REPORT_ROOT = Path(__file__).resolve().parent
RUNS_DIR = REPORT_ROOT / "runs"
ASSETS_DIR = REPORT_ROOT / "assets"

INK = "#17233A"
BLUE = "#2364AA"
TEAL = "#0F8B8D"
GOLD = "#C98A2C"
MUTED = "#5B677A"
GRID = "#D8DEE9"
PALE_BLUE = "#EAF2FB"
PALE_GOLD = "#FFF4DD"


class MissingReportDependency(RuntimeError):
    pass


def require_report_dependencies() -> dict[str, Any]:
    try:
        import fitz
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
        from PIL import Image as PILImage
        from PIL import ImageDraw, ImageFont
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            Image,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ModuleNotFoundError as exc:
        raise MissingReportDependency(
            "Missing report dependency. Install with: "
            "python3 -m pip install -r docs/reports/mlx-cleanup-benchmark/requirements.txt"
        ) from exc

    return locals()


def parse_ollama_report(path: Path) -> dict[str, Any]:
    return parse_ollama_text(path.read_text(encoding="utf-8"))


def parse_ollama_text(text: str) -> dict[str, Any]:
    rows: dict[str, dict[str, Any]] = {}
    median = p90 = first_call = None
    model = "gemma3:4b"

    for line in text.splitlines():
        heading_match = re.fullmatch(r"##\s+(.+)", line.strip())
        if heading_match:
            model = heading_match.group(1).strip()

        if line.startswith("Median "):
            vals = re.findall(r"([0-9]+\.[0-9]+)\s+s", line)
            if len(vals) >= 3:
                median, p90, first_call = map(float, vals[:3])

        if not _is_markdown_table_data_row(line):
            continue
        parts = _table_parts(line)
        if len(parts) >= 3 and re.fullmatch(r"\d+", parts[0]):
            rows[parts[0]] = {
                "latency": float(parts[1].replace(" s", "")),
                "output": parts[2],
            }

    _require_number(median, "Ollama median")
    _require_number(p90, "Ollama p90")
    _require_number(first_call, "Ollama first call")
    return {
        "provider_id": "ollama",
        "model": model,
        "median": median,
        "p90": p90,
        "first_call": first_call,
        "model_footprint": "3.3 GB",
        "rows": rows,
    }


def parse_mlx_report(path: Path) -> dict[str, Any]:
    return parse_mlx_text(path.read_text(encoding="utf-8"))


def parse_mlx_text(text: str) -> dict[str, Any]:
    rows: dict[str, dict[str, Any]] = {}
    summary: dict[str, Any] = {
        "provider_id": "mlx",
        "model": "mlx-community/Qwen2.5-1.5B-Instruct-4bit",
        "model_footprint": "0.88 GB",
    }

    for line in text.splitlines():
        model_match = re.fullmatch(r"- Model:\s+`?([^`]+)`?", line.strip())
        if model_match:
            summary["model"] = model_match.group(1).strip()

        for key, output_key in [
            ("Load", "load"),
            ("Warmup", "warmup"),
            ("Median", "median"),
            ("p90", "p90"),
        ]:
            if line.startswith(f"- {key}:"):
                match = re.search(r"([0-9]+\.[0-9]+)", line)
                if match:
                    summary[output_key] = float(match.group(1))

        if not _is_markdown_table_data_row(line):
            continue
        parts = _table_parts(line)
        if len(parts) >= 4 and re.fullmatch(r"\d+", parts[0]):
            rows[parts[0]] = {
                "latency": float(parts[1].replace(" s", "")),
                "reference": parts[2],
                "output": parts[3],
            }

    for key in ["load", "warmup", "median", "p90"]:
        _require_number(summary.get(key), f"MLX {key}")
    summary["rows"] = rows
    return summary


def _is_markdown_table_data_row(line: str) -> bool:
    return (
        line.startswith("| ")
        and not line.startswith("| id")
        and not line.startswith("| ---")
    )


def _table_parts(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def _require_number(value: Any, label: str) -> None:
    if not isinstance(value, float):
        raise ValueError(f"Missing {label} value in raw benchmark report")


def build_run_folder_name(run_date: str, run_label: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", run_label.lower()).strip("-")
    if not slug:
        raise ValueError("Run label must contain at least one letter or number")
    return f"{run_date}-{slug}"


def create_run_folder(base_dir: Path, run_date: str, run_label: str, force: bool) -> Path:
    run_dir = base_dir / build_run_folder_name(run_date, run_label)
    if run_dir.exists():
        if not force:
            raise FileExistsError(f"Run folder already exists: {run_dir}")
        shutil.rmtree(run_dir)
    (run_dir / "charts").mkdir(parents=True)
    return run_dir


def machine_context() -> dict[str, str]:
    context = {
        "platform": platform.platform(),
        "machine": platform.machine(),
        "processor": platform.processor(),
        "python": platform.python_version(),
    }
    try:
        chip = subprocess.run(
            ["sysctl", "-n", "machdep.cpu.brand_string"],
            text=True,
            capture_output=True,
            check=False,
        ).stdout.strip()
        memory = subprocess.run(
            ["sysctl", "-n", "hw.memsize"],
            text=True,
            capture_output=True,
            check=False,
        ).stdout.strip()
        if chip:
            context["chip"] = chip
        if memory:
            context["memory_bytes"] = memory
    except OSError:
        pass
    return context


def write_structured_outputs(
    run_dir: Path,
    run_label: str,
    ollama: dict[str, Any],
    mlx: dict[str, Any],
) -> None:
    sample_ids = sorted(set(ollama["rows"]).intersection(mlx["rows"]), key=int)
    created_at = datetime.now(timezone.utc).isoformat()
    metadata = {
        "run_label": run_label,
        "created_at": created_at,
        "samples_manifest": "bench/samples/manifest.json",
        "sample_count": len(sample_ids),
        "machine_context": machine_context(),
        "ollama_runtime": "localhost Ollama /api/generate",
        "mlx_runtime": "Python mlx-lm proof output, Swift-native MLX target",
        "notes": [
            "Task 1 formalizes the first MLX vs. Ollama cleanup benchmark report.",
            "Python mlx-lm is proof tooling, not the Local Flow app runtime target.",
        ],
    }
    results = {
        "run_label": run_label,
        "samples": [
            {
                "id": sample_id,
                "reference": mlx["rows"][sample_id].get("reference", ""),
            }
            for sample_id in sample_ids
        ],
        "providers": [
            _provider_results(ollama, sample_ids),
            _provider_results(mlx, sample_ids),
        ],
        "acceptance": {
            "decision": "fail",
            "reason": (
                "Latency target passed, but MLX output still needs sanitizer and "
                "meaning-preservation guards before it can become the default provider."
            ),
            "required_next_controls": [
                "Strip Markdown/code formatting.",
                "Reject added words or added content.",
                "Preserve protected markers and personal dictionary terms.",
                "Fall back to raw transcript when output is unsafe.",
            ],
        },
    }
    (run_dir / "metadata.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    (run_dir / "cleanup-results.json").write_text(json.dumps(results, indent=2) + "\n", encoding="utf-8")


def _provider_results(provider: dict[str, Any], sample_ids: list[str]) -> dict[str, Any]:
    rows = provider["rows"]
    return {
        "provider_id": provider["provider_id"],
        "model": provider["model"],
        "model_footprint": provider["model_footprint"],
        "load_seconds": provider.get("load"),
        "warmup_seconds": provider.get("warmup"),
        "first_call_seconds": provider.get("first_call"),
        "median_seconds": provider["median"],
        "p90_seconds": provider["p90"],
        "per_sample": [
            {
                "id": sample_id,
                "reference": rows[sample_id].get("reference", ""),
                "latency_seconds": rows[sample_id]["latency"],
                "cleaned_output": rows[sample_id]["output"],
                "quality_flags": _quality_flags(provider["provider_id"], rows[sample_id]["output"]),
            }
            for sample_id in sample_ids
        ],
    }


def _quality_flags(provider_id: str, output: str) -> list[str]:
    flags: list[str] = []
    if "`" in output:
        flags.append("markdown_artifact")
    if provider_id == "mlx" and "the Standup" in output:
        flags.append("added_content_risk")
    if not flags:
        flags.append("meaning_review_required")
    return flags


def report_metadata_text(run_label: str) -> str:
    return f"Research brief | Run label: {run_label} | July 6, 2026 | Local benchmark over 24 filler-heavy dictation references"


def build_markdown_report(run_dir: Path, run_label: str, ollama: dict[str, Any], mlx: dict[str, Any]) -> Path:
    speedup_median = ollama["median"] / mlx["median"]
    speedup_p90 = ollama["p90"] / mlx["p90"]
    load_warmup = mlx["load"] + mlx["warmup"]
    lines = [
        "# Local Flow MLX Cleanup Benchmark",
        "",
        f"Run label: `{run_label}`",
        "",
        "## Executive Summary",
        "",
        (
            "The benchmark supports continuing Phase 3 as MLX Performance and "
            "Reliability Hardening. MLX delivered a large cleanup-latency win, "
            "but it is not ready to become the default provider until sanitizer "
            "and meaning-preservation guards are implemented."
        ),
        "",
        "## Performance Metrics",
        "",
        f"- Ollama `{ollama['model']}`: median `{ollama['median']:.3f}s`, p90 `{ollama['p90']:.3f}s`, first call `{ollama['first_call']:.3f}s`.",
        f"- MLX `{mlx['model']}`: load `{mlx['load']:.3f}s`, warmup `{mlx['warmup']:.3f}s`, median `{mlx['median']:.3f}s`, p90 `{mlx['p90']:.3f}s`.",
        f"- Median speedup: `{speedup_median:.1f}x`.",
        f"- p90 speedup: `{speedup_p90:.1f}x`.",
        f"- MLX load + warmup: `{load_warmup:.3f}s`.",
        "",
        "## Model Comparison",
        "",
        "| Dimension | Ollama | MLX |",
        "| --- | --- | --- |",
        f"| Model | `{ollama['model']}` | `{mlx['model']}` |",
        f"| Footprint | `{ollama['model_footprint']}` | `{mlx['model_footprint']}` |",
        "| Runtime posture | Localhost daemon | Native MLX target |",
        "| App-runtime target | Existing fallback | Swift-native MLX, not Python |",
        "",
        "## Acceptance Status",
        "",
        "Acceptance decision: `fail` for making MLX the production default.",
        "",
        "Reason: latency passed the Phase 3 target, but MLX output still needs sanitizer and guard work:",
        "",
        "- Strip Markdown/code formatting.",
        "- Reject added words or added content.",
        "- Preserve protected markers and personal dictionary terms.",
        "- Fall back to raw transcript when output is unsafe.",
        "",
        "## Known Output Findings",
        "",
        "- Markdown/code formatting appeared around `GROUP BY`.",
        "- `standup` became `the Standup`.",
        "- Some filler words were retained.",
        "",
        "## Generated Artifacts",
        "",
        "- `report.pdf`: portfolio-grade visual report.",
        "- `report.docx`: editable Word-compatible report.",
        "- `report.md`: agent-readable summary.",
        "- `metadata.json` and `cleanup-results.json`: structured run data.",
        "- `charts/*.png`: reusable chart assets.",
        "",
        "## Source Inputs",
        "",
        "- `raw-ollama-cleanup.md`",
        "- `raw-mlx-cleanup.md`",
        "- `bench/samples/manifest.json`",
        "",
    ]
    out = run_dir / "report.md"
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def rgb(hex_color: str) -> tuple[int, int, int]:
    value = hex_color.replace("#", "")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def font(deps: dict[str, Any], size: int, bold: bool = False) -> Any:
    ImageFont = deps["ImageFont"]
    preferred = "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf"
    try:
        return ImageFont.truetype(preferred, size)
    except OSError:
        return ImageFont.load_default()


def text_size(draw: Any, text: str, fnt: Any) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def draw_centered(draw: Any, xy: tuple[float, float], text: str, fnt: Any, fill: tuple[int, int, int]) -> None:
    x, y = xy
    w, h = text_size(draw, text, fnt)
    draw.text((x - w / 2, y - h / 2), text, font=fnt, fill=fill)


def draw_axis(
    deps: dict[str, Any],
    draw: Any,
    left: int,
    top: int,
    right: int,
    bottom: int,
    max_value: float,
    y_label: str = "Seconds",
) -> None:
    draw.line((left, bottom, right, bottom), fill=rgb(GRID), width=2)
    draw.line((left, top, left, bottom), fill=rgb(GRID), width=2)
    tick_font = font(deps, 20)
    for index in range(5):
        value = max_value * index / 4
        y = bottom - (bottom - top) * index / 4
        draw.line((left - 8, y, right, y), fill=rgb("#EEF2F7"), width=1)
        draw.text((left - 75, y - 12), f"{value:.1f}", font=tick_font, fill=rgb(MUTED))
    draw.text((left - 90, top - 36), y_label, font=font(deps, 20, True), fill=rgb(MUTED))


def save_latency_summary_chart(charts_dir: Path, ollama: dict[str, Any], mlx: dict[str, Any], deps: dict[str, Any]) -> Path:
    PILImage = deps["PILImage"]
    ImageDraw = deps["ImageDraw"]
    out = charts_dir / "latency-summary.png"
    labels = ["Median", "p90"]
    ollama_vals = [ollama["median"], ollama["p90"]]
    mlx_vals = [mlx["median"], mlx["p90"]]

    img = PILImage.new("RGB", (1300, 720), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Warm cleanup latency: MLX vs. Ollama", font=font(deps, 34, True), fill=rgb(INK))
    draw.text((70, 88), "Lower is better. Values are measured over the same 24 reference transcripts.", font=font(deps, 20), fill=rgb(MUTED))
    left, top, right, bottom = 150, 160, 1210, 610
    draw_axis(deps, draw, left, top, right, bottom, 1.6)
    for center, label, ov, mv in zip([430, 880], labels, ollama_vals, mlx_vals):
        for offset, value, color in [(-58, ov, BLUE), (58, mv, TEAL)]:
            x0 = center + offset - 48
            x1 = center + offset + 48
            y0 = bottom - (value / 1.6) * (bottom - top)
            draw.rounded_rectangle((x0, y0, x1, bottom), radius=10, fill=rgb(color))
            draw_centered(draw, (center + offset, y0 - 22), f"{value:.3f}s", font(deps, 20, True), rgb(INK))
        draw_centered(draw, (center, bottom + 42), label, font(deps, 24, True), rgb(INK))
    draw.rectangle((884, 70, 908, 94), fill=rgb(BLUE))
    draw.text((918, 66), "Ollama gemma3:4b", font=font(deps, 18), fill=rgb(INK))
    draw.rectangle((884, 101, 908, 125), fill=rgb(TEAL))
    draw.text((918, 97), "MLX Qwen2.5 1.5B", font=font(deps, 18), fill=rgb(INK))
    img.save(out)
    return out


def save_per_sample_chart(charts_dir: Path, ollama: dict[str, Any], mlx: dict[str, Any], deps: dict[str, Any]) -> Path:
    PILImage = deps["PILImage"]
    ImageDraw = deps["ImageDraw"]
    out = charts_dir / "per-sample-latency.png"
    ids = sorted(set(ollama["rows"]).intersection(mlx["rows"]), key=int)
    ollama_vals = [ollama["rows"][sample_id]["latency"] for sample_id in ids]
    mlx_vals = [mlx["rows"][sample_id]["latency"] for sample_id in ids]

    img = PILImage.new("RGB", (1400, 740), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Per-utterance cleanup latency", font=font(deps, 34, True), fill=rgb(INK))
    draw.text((70, 88), "Same 24 references, same task. The MLX path stays consistently below the Ollama baseline.", font=font(deps, 20), fill=rgb(MUTED))
    left, top, right, bottom = 120, 160, 1320, 610
    draw_axis(deps, draw, left, top, right, bottom, 1.8)

    def point(index: int, value: float) -> tuple[float, float]:
        x = left + (right - left) * index / (len(ids) - 1)
        y = bottom - (value / 1.8) * (bottom - top)
        return x, y

    for values, color in [(ollama_vals, BLUE), (mlx_vals, TEAL)]:
        pts = [point(index, value) for index, value in enumerate(values)]
        draw.line(pts, fill=rgb(color), width=5, joint="curve")
        for x, y in pts:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=rgb(color))
    for index, sample_id in enumerate(ids):
        if index % 2 == 0:
            x, _ = point(index, 0)
            draw_centered(draw, (x, bottom + 32), sample_id, font(deps, 16), rgb(MUTED))
    draw.line((996, 54, 1030, 54), fill=rgb(BLUE), width=5)
    draw.text((1040, 42), "Ollama gemma3:4b", font=font(deps, 19), fill=rgb(INK))
    draw.line((996, 88, 1030, 88), fill=rgb(TEAL), width=5)
    draw.text((1040, 76), "MLX Qwen2.5 1.5B", font=font(deps, 19), fill=rgb(INK))
    img.save(out)
    return out


def save_model_footprint_chart(charts_dir: Path, deps: dict[str, Any]) -> Path:
    PILImage = deps["PILImage"]
    ImageDraw = deps["ImageDraw"]
    out = charts_dir / "model-footprint.png"
    img = PILImage.new("RGB", (1150, 640), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Runtime model footprint", font=font(deps, 34, True), fill=rgb(INK))
    draw.text((70, 88), "Smaller resident model improves the odds of fast cleanup under normal desktop load.", font=font(deps, 20), fill=rgb(MUTED))
    left, top, right, bottom = 150, 155, 1060, 540
    draw_axis(deps, draw, left, top, right, bottom, 3.6, "GB")
    for center, label, value, color in [
        (430, ["Ollama", "gemma3:4b"], 3.3, BLUE),
        (780, ["MLX", "Qwen2.5 1.5B"], 0.88, TEAL),
    ]:
        y0 = bottom - (value / 3.6) * (bottom - top)
        draw.rounded_rectangle((center - 85, y0, center + 85, bottom), radius=12, fill=rgb(color))
        draw_centered(draw, (center, y0 - 24), f"{value:.2f} GB", font(deps, 22, True), rgb(INK))
        for row, part in enumerate(label):
            draw_centered(draw, (center, bottom + 34 + row * 24), part, font(deps, 20, row == 0), rgb(INK))
    img.save(out)
    return out


def save_pipeline_visual(charts_dir: Path, deps: dict[str, Any]) -> Path:
    PILImage = deps["PILImage"]
    ImageDraw = deps["ImageDraw"]
    out = charts_dir / "pipeline-comparison.png"
    img = PILImage.new("RGB", (1400, 520), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "Cleanup architecture shift", font=font(deps, 34, True), fill=rgb(INK))

    def box(x: int, y: int, w: int, h: int, text: str, fill: str) -> None:
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=rgb(fill), outline=rgb("#FFFFFF"), width=3)
        lines = text.split("\n")
        for index, line in enumerate(lines):
            draw_centered(draw, (x + w / 2, y + h / 2 - (len(lines) - 1) * 12 + index * 24), line, font(deps, 19, index == 0), rgb(INK))

    def arrow(x1: int, y: int, x2: int) -> None:
        draw.line((x1, y, x2, y), fill=rgb(MUTED), width=4)
        draw.polygon([(x2, y), (x2 - 16, y - 9), (x2 - 16, y + 9)], fill=rgb(MUTED))

    draw.text((70, 112), "Current baseline", font=font(deps, 22, True), fill=rgb(BLUE))
    y = 150
    box(70, y, 220, 90, "Raw\ntranscript", PALE_BLUE)
    arrow(300, y + 45, 365)
    box(370, y, 260, 90, "Ollama HTTP\nlocalhost:11434", "#DCE8F7")
    arrow(640, y + 45, 705)
    box(710, y, 250, 90, "gemma3:4b\n3.3 GB", "#DCE8F7")
    arrow(970, y + 45, 1035)
    box(1040, y, 220, 90, "Cleaned\ntext", PALE_BLUE)

    draw.text((70, 292), "Proposed hardening target", font=font(deps, 22, True), fill=rgb(TEAL))
    y = 330
    box(70, y, 220, 90, "Raw\ntranscript", "#E1F4F2")
    arrow(300, y + 45, 365)
    box(370, y, 260, 90, "MLX runtime\nApple Silicon", "#D9F0EE")
    arrow(640, y + 45, 705)
    box(710, y, 250, 90, "Qwen2.5 1.5B\n0.88 GB", "#D9F0EE")
    arrow(970, y + 45, 1035)
    box(1040, y, 220, 90, "Cleaned\ntext", "#E1F4F2")
    img.save(out)
    return out


def generate_charts(run_dir: Path, ollama: dict[str, Any], mlx: dict[str, Any], deps: dict[str, Any]) -> dict[str, Path]:
    charts_dir = run_dir / "charts"
    return {
        "latency": save_latency_summary_chart(charts_dir, ollama, mlx, deps),
        "per_sample": save_per_sample_chart(charts_dir, ollama, mlx, deps),
        "footprint": save_model_footprint_chart(charts_dir, deps),
        "pipeline": save_pipeline_visual(charts_dir, deps),
    }


def set_cell_shading(cell: Any, fill: str, deps: dict[str, Any]) -> None:
    OxmlElement = deps["OxmlElement"]
    qn = deps["qn"]
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell: Any, text: str, deps: dict[str, Any], bold: bool = False, color: str = INK, size: float = 9) -> None:
    Pt = deps["Pt"]
    RGBColor = deps["RGBColor"]
    WD_CELL_VERTICAL_ALIGNMENT = deps["WD_CELL_VERTICAL_ALIGNMENT"]
    qn = deps["qn"]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table: Any, deps: dict[str, Any], color: str = "D8DEE9") -> None:
    OxmlElement = deps["OxmlElement"]
    qn = deps["qn"]
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_docx_para(doc: Any, deps: dict[str, Any], text: str, size: float = 10.5, bold: bool = False, color: str = INK, after: int = 6, align: Any = None) -> Any:
    Pt = deps["Pt"]
    RGBColor = deps["RGBColor"]
    qn = deps["qn"]
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    return paragraph


def style_docx(doc: Any, deps: dict[str, Any]) -> None:
    Inches = deps["Inches"]
    Pt = deps["Pt"]
    RGBColor = deps["RGBColor"]
    WD_ALIGN_PARAGRAPH = deps["WD_ALIGN_PARAGRAPH"]
    qn = deps["qn"]
    section = doc.sections[0]
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, "#1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "Local Flow MLX Cleanup Benchmark"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED.replace("#", ""))

    footer = section.footer.paragraphs[0]
    footer.text = "Research brief - generated from local benchmark artifacts"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED.replace("#", ""))


def add_metric_table_docx(doc: Any, deps: dict[str, Any], metrics: list[tuple[str, str, str, str]]) -> None:
    WD_TABLE_ALIGNMENT = deps["WD_TABLE_ALIGNMENT"]
    WD_ALIGN_PARAGRAPH = deps["WD_ALIGN_PARAGRAPH"]
    Pt = deps["Pt"]
    RGBColor = deps["RGBColor"]
    table = doc.add_table(rows=1, cols=len(metrics))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, deps, "FFFFFF")
    for index, (label, value, detail, fill) in enumerate(metrics):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, fill, deps)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor.from_string(INK.replace("#", ""))
        paragraph.add_run("\n")
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED.replace("#", ""))
        paragraph.add_run("\n")
        run = paragraph.add_run(detail)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED.replace("#", ""))


def add_comparison_table_docx(doc: Any, deps: dict[str, Any], rows: list[tuple[str, str, str, str]]) -> None:
    WD_TABLE_ALIGNMENT = deps["WD_TABLE_ALIGNMENT"]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, ["Dimension", "Ollama gemma3:4b", "MLX Qwen2.5 1.5B", "Interpretation"]):
        set_cell_shading(cell, "E8EEF5", deps)
        set_cell_text(cell, header, deps, bold=True, size=8.8)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text, deps, size=8.5)
    set_table_borders(table, deps)


def build_docx(run_dir: Path, run_label: str, charts: dict[str, Path], ollama: dict[str, Any], mlx: dict[str, Any], deps: dict[str, Any]) -> Path:
    Document = deps["Document"]
    Inches = deps["Inches"]
    WD_ALIGN_PARAGRAPH = deps["WD_ALIGN_PARAGRAPH"]
    doc = Document()
    style_docx(doc, deps)

    add_docx_para(doc, deps, "LOCAL FLOW", size=10, bold=True, color=TEAL, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_docx_para(doc, deps, "MLX Cleanup Benchmark", size=26, bold=True, color=INK, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_docx_para(doc, deps, "Replacing a 3.3 GB Ollama cleanup hop with a lightweight native MLX model", size=12.5, color=MUTED, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_docx_para(doc, deps, report_metadata_text(run_label), size=9.5, color=MUTED, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_metric_table_docx(doc, deps, [
        ("median cleanup latency", "0.240s", "MLX rerun after unloading 30B model", "E1F4F2"),
        ("median speedup", "3.8x", "0.914s Ollama / 0.240s MLX", "EAF2FB"),
        ("p90 speedup", "5.6x", "1.485s Ollama / 0.264s MLX", "FFF4DD"),
    ])

    doc.add_heading("Executive Summary", level=1)
    add_docx_para(doc, deps, "The benchmark supports making MLX cleanup the focus of the next Local Flow hardening phase. With Qwen2.5 1.5B Instruct 4-bit loaded through MLX, the cleanup hop dropped from a 0.914 second Ollama median to 0.240 seconds, while keeping the dictation path local and avoiding cloud inference.")
    add_docx_para(doc, deps, "This is not yet a production swap. The MLX path needs a deterministic sanitizer, an added-word guard, and an acceptance suite before it replaces Ollama.", bold=True)
    doc.add_picture(str(charts["latency"]), width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Research Method", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = deps["WD_TABLE_ALIGNMENT"].CENTER
    for cell, header in zip(table.rows[0].cells, ["Method item", "Value"]):
        set_cell_shading(cell, "E8EEF5", deps)
        set_cell_text(cell, header, deps, bold=True, size=9)
    for label, value in [
        ("Input set", "24 verbatim references from bench/samples/manifest.json"),
        ("Baseline", "Existing Ollama gemma3:4b cleanup report"),
        ("Candidate", "mlx-community/Qwen2.5-1.5B-Instruct-4bit"),
        ("Sampling", "Temperature 0.0 for the strict rerun"),
        ("Runtime condition", "Qwen3-Coder 30B was unloaded before the final rerun to reduce memory pressure"),
        ("Decision standard", "Lower latency is useful only if meaning preservation remains enforceable"),
    ]:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, deps, bold=True, size=8.6)
        set_cell_text(cells[1], value, deps, size=8.6)
    set_table_borders(table, deps)

    doc.add_heading("Measured Results", level=1)
    add_comparison_table_docx(doc, deps, [
        ("Median warm cleanup", f"{ollama['median']:.3f}s", f"{mlx['median']:.3f}s", "MLX is 3.8x faster"),
        ("p90 warm cleanup", f"{ollama['p90']:.3f}s", f"{mlx['p90']:.3f}s", "MLX is 5.6x faster at the tail"),
        ("First-call / load effect", f"{ollama['first_call']:.3f}s first call", f"{mlx['load'] + mlx['warmup']:.3f}s load + warmup", "MLX has lower observed residency penalty"),
        ("Model footprint", "3.3 GB", "0.88 GB cached files", "MLX candidate is materially lighter"),
        ("Runtime posture", "Localhost daemon", "Native MLX target", "Both stay local; MLX better matches Mac-native goal"),
    ])
    doc.add_picture(str(charts["per_sample"]), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(charts["footprint"]), width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Architecture Implication", level=1)
    doc.add_picture(str(charts["pipeline"]), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_para(doc, deps, "The product goal is not simply local cleanup; it is fast local cleanup. The MLX path removes the Ollama daemon hop, uses a smaller model, and aligns the cleanup engine with Apple Silicon.")

    doc.add_heading("Quality Findings", level=1)
    add_comparison_table_docx(doc, deps, [
        ("Meaning preservation", "No obvious answer-style hallucinations in the strict MLX rerun.", "", ""),
        ("Known artifacts", "One Markdown artifact around GROUP BY; one added article in 'the Standup'; a few filler words retained.", "", ""),
        ("Production guard", "Strip Markdown/code formatting, reject added content, preserve protected markers, and fall back to raw.", "", ""),
        ("Acceptance gate", "Run the 24-reference set plus the production plausibility guard before making MLX default.", "", ""),
    ])

    doc.add_heading("Recommendation", level=1)
    add_docx_para(doc, deps, "Phase 3 Continuation should be scoped as MLX Performance and Reliability Hardening. The first task is not a blind replacement of Ollama; it is an optional MLX cleanup provider with deterministic output controls, benchmark reporting, and rollback to raw or Ollama when safety checks fail.")
    add_docx_para(doc, deps, "Recommended acceptance: MLX median cleanup below 0.350 seconds, p90 below 0.500 seconds, zero meaning changes on the acceptance set, no Markdown artifacts, and no non-local network dependency in the dictation path.", bold=True, color=TEAL)

    doc.add_heading("Sources And Reproducibility", level=1)
    for source in [
        "bench/samples/manifest.json - 24 reference transcripts used as benchmark inputs",
        "raw-ollama-cleanup.md - Ollama gemma3:4b baseline report",
        "raw-mlx-cleanup.md - final MLX rerun report",
        "app/Sources/LocalFlow/Cleaner.swift - production cleanup prompt and plausibility-guard context",
        "https://huggingface.co/mlx-community/Qwen2.5-1.5B-Instruct-4bit - approved MLX model package",
        "https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct - upstream model card",
    ]:
        add_docx_para(doc, deps, source, size=8.8, color=MUTED, after=3)

    out = run_dir / "report.docx"
    doc.save(out)
    return out


def make_pdf_styles(deps: dict[str, Any]) -> Any:
    colors = deps["colors"]
    ParagraphStyle = deps["ParagraphStyle"]
    TA_CENTER = deps["TA_CENTER"]
    getSampleStyleSheet = deps["getSampleStyleSheet"]
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverKicker", fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor(TEAL), alignment=TA_CENTER, spaceAfter=8))
    styles.add(ParagraphStyle(name="CoverTitle", fontName="Helvetica-Bold", fontSize=28, textColor=colors.HexColor(INK), alignment=TA_CENTER, leading=32, spaceAfter=6))
    styles.add(ParagraphStyle(name="CoverSub", fontName="Helvetica", fontSize=12, textColor=colors.HexColor(MUTED), alignment=TA_CENTER, leading=16, spaceAfter=14))
    styles.add(ParagraphStyle(name="H1Custom", fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor(BLUE), spaceBefore=14, spaceAfter=7))
    styles.add(ParagraphStyle(name="BodyCustom", fontName="Helvetica", fontSize=9.5, textColor=colors.HexColor(INK), leading=13, spaceAfter=7))
    styles.add(ParagraphStyle(name="SmallCustom", fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor(MUTED), leading=9, spaceAfter=3))
    styles.add(ParagraphStyle(name="Callout", fontName="Helvetica-Bold", fontSize=10, textColor=colors.HexColor(INK), leading=14, backColor=colors.HexColor(PALE_GOLD), borderPadding=8, spaceAfter=8))
    return styles


def pdf_metric_table(deps: dict[str, Any], metrics: list[tuple[str, str, str, str]]) -> Any:
    colors = deps["colors"]
    inch = deps["inch"]
    Paragraph = deps["Paragraph"]
    Table = deps["Table"]
    TableStyle = deps["TableStyle"]
    getSampleStyleSheet = deps["getSampleStyleSheet"]
    rows = [
        Paragraph(f"<b><font size='18'>{value}</font></b><br/><font color='{MUTED}'><b>{label}</b><br/>{detail}</font>", getSampleStyleSheet()["BodyText"])
        for label, value, detail, _fill in metrics
    ]
    table = Table([rows], colWidths=[2.05 * inch, 2.05 * inch, 2.05 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#E1F4F2")),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#EAF2FB")),
        ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#FFF4DD")),
        ("INNERGRID", (0, 0), (-1, -1), 2, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]))
    return table


def build_pdf(run_dir: Path, run_label: str, charts: dict[str, Path], ollama: dict[str, Any], mlx: dict[str, Any], deps: dict[str, Any]) -> Path:
    colors = deps["colors"]
    letter = deps["letter"]
    inch = deps["inch"]
    Image = deps["Image"]
    PageBreak = deps["PageBreak"]
    Paragraph = deps["Paragraph"]
    SimpleDocTemplate = deps["SimpleDocTemplate"]
    Spacer = deps["Spacer"]
    Table = deps["Table"]
    TableStyle = deps["TableStyle"]
    styles = make_pdf_styles(deps)
    out = run_dir / "report.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=letter, rightMargin=0.65 * inch, leftMargin=0.65 * inch, topMargin=0.65 * inch, bottomMargin=0.62 * inch)
    story: list[Any] = []
    story.append(Spacer(1, 0.45 * inch))
    story.append(Paragraph("LOCAL FLOW", styles["CoverKicker"]))
    story.append(Paragraph("MLX Cleanup Benchmark", styles["CoverTitle"]))
    story.append(Paragraph("Replacing a 3.3 GB Ollama cleanup hop with a lightweight native MLX model", styles["CoverSub"]))
    story.append(Paragraph(report_metadata_text(run_label), styles["SmallCustom"]))
    story.append(Spacer(1, 0.18 * inch))
    story.append(pdf_metric_table(deps, [
        ("median cleanup latency", "0.240s", "MLX after unloading 30B model", "E1F4F2"),
        ("median speedup", "3.8x", "0.914s / 0.240s", "EAF2FB"),
        ("p90 speedup", "5.6x", "1.485s / 0.264s", "FFF4DD"),
    ]))
    story.append(Spacer(1, 0.18 * inch))
    story.append(Image(str(charts["latency"]), width=6.4 * inch, height=3.45 * inch))
    story.append(PageBreak())

    story.append(Paragraph("Executive Summary", styles["H1Custom"]))
    story.append(Paragraph("The benchmark supports making MLX cleanup the center of the next Local Flow hardening phase. With Qwen2.5 1.5B Instruct 4-bit loaded through MLX, the cleanup hop dropped from a 0.914 second Ollama median to 0.240 seconds while keeping the dictation path local.", styles["BodyCustom"]))
    story.append(Paragraph("This is a performance proof, not a blind production swap. The MLX path still needs deterministic output sanitization and guardrails before becoming the default.", styles["Callout"]))
    story.append(Paragraph("Research Method", styles["H1Custom"]))
    method = [
        ["Input set", "24 references from bench/samples/manifest.json"],
        ["Baseline", "Ollama gemma3:4b cleanup report"],
        ["Candidate", "mlx-community/Qwen2.5-1.5B-Instruct-4bit"],
        ["Sampling", "Temperature 0.0 strict rerun"],
        ["Runtime condition", "Qwen3-Coder 30B unloaded before final rerun"],
        ["Decision standard", "Lower latency only counts if meaning preservation remains enforceable"],
    ]
    table = Table([["Method item", "Value"]] + method, colWidths=[1.55 * inch, 4.8 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.3),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(PageBreak())

    story.append(Paragraph("Measured Results", styles["H1Custom"]))
    comparison = [
        ["Dimension", "Ollama gemma3:4b", "MLX Qwen2.5 1.5B", "Interpretation"],
        ["Median warm cleanup", f"{ollama['median']:.3f}s", f"{mlx['median']:.3f}s", "3.8x faster"],
        ["p90 warm cleanup", f"{ollama['p90']:.3f}s", f"{mlx['p90']:.3f}s", "5.6x faster at the tail"],
        ["First-call/load effect", f"{ollama['first_call']:.3f}s first call", f"{mlx['load'] + mlx['warmup']:.3f}s load + warmup", "Lower residency penalty"],
        ["Model footprint", "3.3 GB", "0.88 GB cached", "Materially lighter"],
    ]
    table = Table(comparison, colWidths=[1.55 * inch, 1.5 * inch, 1.55 * inch, 1.75 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.8),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(Spacer(1, 0.12 * inch))
    story.append(Image(str(charts["per_sample"]), width=6.45 * inch, height=3.3 * inch))
    story.append(Spacer(1, 0.08 * inch))
    story.append(Image(str(charts["footprint"]), width=5.65 * inch, height=3.08 * inch))
    story.append(PageBreak())

    story.append(Paragraph("Architecture Implication", styles["H1Custom"]))
    story.append(Image(str(charts["pipeline"]), width=6.45 * inch, height=2.42 * inch))
    story.append(Paragraph("The product goal is fast local cleanup. MLX removes the Ollama daemon hop, uses a smaller model, and aligns the cleanup engine with Apple Silicon.", styles["BodyCustom"]))
    story.append(Paragraph("Quality Findings", styles["H1Custom"]))
    findings = [
        ["Finding", "Implication"],
        ["Meaning preservation", "No obvious answer-style hallucinations in the strict MLX rerun."],
        ["Known artifacts", "Backticks around GROUP BY; one added article in 'the Standup'; some filler retained."],
        ["Production guard", "Strip Markdown/code formatting, reject added content, preserve protected markers, fall back to raw."],
        ["Acceptance gate", "24-reference acceptance set, production plausibility guard, and no non-local network dependency."],
    ]
    table = Table(findings, colWidths=[1.55 * inch, 4.8 * inch])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.0),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(GRID)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(Paragraph("Recommendation", styles["H1Custom"]))
    story.append(Paragraph("Scope Phase 3 Continuation as MLX Performance and Reliability Hardening: implement an optional MLX cleanup provider, deterministic sanitizer, benchmark reporting, and rollback to raw or Ollama when safety checks fail.", styles["BodyCustom"]))
    story.append(Paragraph("Acceptance target: MLX median below 0.350 seconds, p90 below 0.500 seconds, zero meaning changes, no Markdown artifacts, and no non-local network dependency in the dictation path.", styles["Callout"]))
    story.append(Paragraph("Sources", styles["H1Custom"]))
    for source in [
        "bench/samples/manifest.json",
        "raw-ollama-cleanup.md",
        "raw-mlx-cleanup.md",
        "app/Sources/LocalFlow/Cleaner.swift",
        "https://huggingface.co/mlx-community/Qwen2.5-1.5B-Instruct-4bit",
        "https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct",
    ]:
        story.append(Paragraph(source, styles["SmallCustom"]))

    def header_footer(canvas: Any, doc_obj: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(0.65 * inch, 0.35 * inch, "Local Flow MLX Cleanup Benchmark")
        canvas.drawRightString(7.85 * inch, 0.35 * inch, f"Page {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return out


def render_pdf_pages(run_dir: Path, pdf_path: Path, deps: dict[str, Any]) -> tuple[list[Path], Path]:
    fitz = deps["fitz"]
    PILImage = deps["PILImage"]
    render_dir = run_dir / "qa-render"
    render_dir.mkdir(exist_ok=True)
    doc = fitz.open(str(pdf_path))
    page_paths = []
    for index, page in enumerate(doc, start=1):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        out = render_dir / f"pdf-page-{index}.png"
        pix.save(str(out))
        page_paths.append(out)
    doc.close()

    thumbs = []
    for path in page_paths:
        img = PILImage.open(path).convert("RGB")
        img.thumbnail((420, 560))
        thumbs.append(img.copy())
    sheet = PILImage.new("RGB", (920, ((len(thumbs) + 1) // 2) * 610), "white")
    for index, thumb in enumerate(thumbs):
        x = (index % 2) * 460 + 20
        y = (index // 2) * 610 + 20
        sheet.paste(thumb, (x, y))
    contact = render_dir / "pdf-contact-sheet.png"
    sheet.save(contact)
    return page_paths, contact


def generate_report(
    run_label: str,
    run_date: str,
    ollama_raw: Path,
    mlx_raw: Path,
    force: bool,
) -> Path:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    run_dir = create_run_folder(RUNS_DIR, run_date, run_label, force)
    ollama_dest = run_dir / "raw-ollama-cleanup.md"
    mlx_dest = run_dir / "raw-mlx-cleanup.md"
    shutil.copyfile(ollama_raw, ollama_dest)
    shutil.copyfile(mlx_raw, mlx_dest)

    ollama = parse_ollama_report(ollama_dest)
    mlx = parse_mlx_report(mlx_dest)
    write_structured_outputs(run_dir, run_label, ollama, mlx)

    markdown_path = build_markdown_report(run_dir, run_label, ollama, mlx)
    deps = require_report_dependencies()
    charts = generate_charts(run_dir, ollama, mlx, deps)
    docx_path = build_docx(run_dir, run_label, charts, ollama, mlx, deps)
    pdf_path = build_pdf(run_dir, run_label, charts, ollama, mlx, deps)
    pages, contact = render_pdf_pages(run_dir, pdf_path, deps)

    print(f"Run folder: {run_dir}")
    print(f"Markdown: {markdown_path}")
    print(f"DOCX: {docx_path}")
    print(f"PDF: {pdf_path}")
    print(f"PDF pages rendered: {len(pages)}")
    print(f"PDF contact sheet: {contact}")
    return run_dir


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a formal Local Flow MLX cleanup benchmark report.")
    parser.add_argument("--run-label", default="initial-mlx-vs-ollama", help="Human-readable run label.")
    parser.add_argument("--run-date", default="2026-07-06", help="Run date in YYYY-MM-DD form.")
    parser.add_argument("--ollama-raw", type=Path, required=True, help="Path to raw Ollama cleanup Markdown.")
    parser.add_argument("--mlx-raw", type=Path, required=True, help="Path to raw MLX cleanup Markdown.")
    parser.add_argument("--force", action="store_true", help="Overwrite an existing run folder with the same date and label.")
    args = parser.parse_args()
    generate_report(args.run_label, args.run_date, args.ollama_raw, args.mlx_raw, args.force)


if __name__ == "__main__":
    main()
